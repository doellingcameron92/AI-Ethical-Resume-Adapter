"""Regeneration: ``GeneralizedProfile`` -> one uniform, skills-first blinded document.

The template is fixed: every candidate gets the same sections in the same
order, and a section with no content says so explicitly instead of leaving a
redaction marker. ``RenderedDocument`` lines keep their source spans so the
audit can verify them before anything is written.
"""

from __future__ import annotations

import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from .models import GeneralizedProfile, RenderedDocument, RenderedLine

SECTIONS: tuple[str, ...] = ("Skills", "Certifications", "Experience", "Education", "Achievements")
TITLE_PREFIX = "Blinded Resume"
EMPTY_SECTION = "None listed"
FOOTER = ("Standardized, de-identified document for human review. It is not a score, ranking, "
          "or hiring recommendation.")
TITLE_RE = re.compile(rf"^{TITLE_PREFIX} · CR-[0-9A-F]{{8}}$")


def is_template_text(text: str) -> bool:
    return text in SECTIONS or text in (EMPTY_SECTION, FOOTER) or bool(TITLE_RE.match(text))


def render(profile: GeneralizedProfile) -> RenderedDocument:
    doc = RenderedDocument(candidate_ref=profile.candidate_ref, as_of=profile.as_of, granularity=profile.granularity)
    add = doc.lines.append
    add(RenderedLine(text=f"{TITLE_PREFIX} · {profile.candidate_ref}", style="title"))
    if profile.home_region:
        add(RenderedLine(text=f"Region: {profile.home_region.text}", style="meta", section="header",
                         derivation="generalized", sources=profile.home_region.sources))

    def section(name: str, items: list[RenderedLine]) -> None:
        add(RenderedLine(text=name, style="section", section=name))
        doc.lines.extend(items or [RenderedLine(text=EMPTY_SECTION, style="detail", section=name)])

    section("Skills", [RenderedLine(text=s.text, style="item", section="Skills", derivation=s.derivation, sources=s.sources)
                       for s in profile.skills])
    section("Certifications", [RenderedLine(text=c.text, style="bullet", section="Certifications", derivation=c.derivation,
                                            sources=c.sources) for c in profile.certifications])
    exp: list[RenderedLine] = []
    for role in profile.experience:
        exp.append(RenderedLine(text=role.title.text, style="entry", section="Experience", derivation=role.title.derivation,
                                sources=role.title.sources, computed=role.duration if role.kind == "career_break" else None))
        if role.context:
            exp.append(RenderedLine(text=role.context.text, style="meta", section="Experience", derivation="generalized",
                                    sources=role.context.sources))
        if role.duration:
            spans = [role.duration.start] + ([role.duration.end] if role.duration.end else [])
            exp.append(RenderedLine(text=f"Duration: {role.duration.text}", style="meta", section="Experience",
                                    derivation="computed", sources=spans, computed=role.duration))
        exp.extend(RenderedLine(text=b.text, style="bullet", section="Experience", derivation=b.derivation, sources=b.sources)
                   for b in role.bullets)
    section("Experience", exp)
    section("Education", [RenderedLine(text=e.text, style="bullet", section="Education", derivation=e.derivation,
                                       sources=e.sources) for e in profile.education])
    section("Achievements", [RenderedLine(text=a.text, style="bullet", section="Achievements", derivation=a.derivation,
                                          sources=a.sources) for a in profile.achievements])
    add(RenderedLine(text=FOOTER, style="footer"))
    return doc


def _blocks(doc: RenderedDocument) -> list[tuple[str, str]]:
    """Group consecutive skill items into one inline block; everything else is one block per line."""
    blocks: list[tuple[str, str]] = []
    for line in doc.lines:
        if line.style == "item" and blocks and blocks[-1][0] == "item":
            blocks[-1] = ("item", blocks[-1][1] + " · " + line.text)
        else:
            blocks.append((line.style, line.text))
    return blocks


def to_markdown(doc: RenderedDocument) -> str:
    prefix = {"title": "# ", "section": "## ", "entry": "### "}
    out: list[str] = []
    prev = ""
    for style, text in _blocks(doc):
        if style == "bullet":
            if prev and prev != "bullet":
                out.append("")
            out.append(f"- {text}")
        else:
            if out:
                out.append("")
            if style == "footer":
                out += ["---", ""]
            if style in ("meta", "footer"):
                out.append(f"_{text}_")
            else:
                out.append(prefix.get(style, "") + text)
        prev = style
    return "\n".join(out) + "\n"


def write_markdown(doc: RenderedDocument, path: Path) -> Path:
    path.write_text(to_markdown(doc), encoding="utf-8")
    return path


def write_docx(doc: RenderedDocument, path: Path) -> Path:
    d = Document()
    d.core_properties.author = ""
    d.core_properties.title = doc.lines[0].text if doc.lines else TITLE_PREFIX
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    for style, text in _blocks(doc):
        if style == "title":
            d.add_heading(text, level=0)
        elif style == "section":
            d.add_heading(text, level=1)
        elif style == "entry":
            d.add_heading(text, level=2)
        elif style == "bullet":
            d.add_paragraph(text, style="List Bullet")
        elif style in ("meta", "footer"):
            d.add_paragraph().add_run(text).italic = True
        else:
            d.add_paragraph(text)
    d.save(str(path))
    return path


def write_pdf(doc: RenderedDocument, path: Path) -> Path:
    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("t", parent=base["Title"], fontSize=16, spaceAfter=6),
        "section": ParagraphStyle("s", parent=base["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4),
        "entry": ParagraphStyle("e", parent=base["Heading4"], fontSize=10.5, spaceBefore=6, spaceAfter=1),
        "meta": ParagraphStyle("m", parent=base["Italic"], fontSize=9, textColor="#444444"),
        "bullet": ParagraphStyle("b", parent=base["BodyText"], fontSize=10, leftIndent=12, bulletIndent=2),
        "detail": ParagraphStyle("d", parent=base["BodyText"], fontSize=10),
        "item": ParagraphStyle("i", parent=base["BodyText"], fontSize=10),
        "footer": ParagraphStyle("f", parent=base["Italic"], fontSize=8, textColor="#666666"),
    }
    story: list = []
    for style, text in _blocks(doc):
        if style == "footer":
            story += [Spacer(1, 10), HRFlowable(width="100%", thickness=0.5), Spacer(1, 4)]
        para = Paragraph(escape(text), styles[style], bulletText="•" if style == "bullet" else None)
        story.append(para)
    pdf = SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch, title=doc.lines[0].text if doc.lines else "",
                            author="", subject="", creator="resume_adapter")
    pdf.build(story)
    return path


WRITERS = {"md": write_markdown, "docx": write_docx, "pdf": write_pdf}
